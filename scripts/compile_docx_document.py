# -*- coding: utf-8 -*-
"""
Compile TMT_OPERATIONS_PLATFORM_V1.1_MASTER.md into
'TMT_Operations_Platform_Intern_Implementation_Ready_Documentation_v1.0 (1).docx'
Preserving original document themes, fonts, colors, styles, callouts, tables, and code blocks.
"""

import os
import sys
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

MD_PATH = os.path.join("docs", "TMT_OPERATIONS_PLATFORM_V1.1_MASTER.md")
TEMPLATE_PATH = "TMT_Operations_Platform_Intern_Implementation_Ready_Documentation_v1.0_ORIGINAL.docx"
OUTPUT_PATH = "TMT_Operations_Platform_Intern_Implementation_Ready_Documentation_v1.0 (1).docx"

print(f"Reading markdown from: {MD_PATH}")
with open(MD_PATH, "r", encoding="utf-8") as f:
    md_content = f.read()

# Load docx template to inherit margins, headers, footers, styles
doc = Document(TEMPLATE_PATH)

# Clear existing elements except sectPr
body = doc._body._element
for el in list(body):
    if el.tag.endswith('sectPr'):
        continue
    body.remove(el)

# Styling helpers
def set_cell_background(cell, fill_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=80, bottom=80, start=100, end=100):
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{start}" w:type="dxa"/><w:right w:w="{end}" w:type="dxa"/></w:tcMar>')
    cell._tc.get_or_add_tcPr().append(tcMar)

def set_cell_valign(cell, valign="center"):
    valign_el = parse_xml(f'<w:vAlign {nsdecls("w")} w:val="{valign}"/>')
    cell._tc.get_or_add_tcPr().append(valign_el)

def add_title_block():
    p0 = doc.add_paragraph(style='Normal')
    p0.paragraph_format.space_after = Pt(2)
    r0 = p0.add_run("TECHMECHA TORQUE\nPRIVATE LIMITED")
    r0.bold = True
    r0.font.size = Pt(18)
    r0.font.color.rgb = RGBColor(0x17, 0x36, 0x5D)

    p1 = doc.add_paragraph(style='Normal')
    p1.paragraph_format.space_before = Pt(4)
    p1.paragraph_format.space_after = Pt(2)
    r1 = p1.add_run("TMT OPERATIONS PLATFORM")
    r1.bold = True
    r1.font.size = Pt(26)
    r1.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)

    p2 = doc.add_paragraph(style='Normal')
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run("Intern Implementation-Ready Documentation")
    r2.bold = True
    r2.font.size = Pt(16)
    r2.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    p3 = doc.add_paragraph(style='Normal')
    p3.paragraph_format.space_after = Pt(8)
    r3 = p3.add_run("Version 1.1 Enterprise Enhanced  |  17 September 2026")
    r3.font.size = Pt(11)
    r3.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    p4 = doc.add_paragraph(style='Normal')
    p4.paragraph_format.space_after = Pt(2)
    r4_title = p4.add_run("Purpose\n")
    r4_title.bold = True
    r4_title.font.size = Pt(11)
    r4_title.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
    r4_body = p4.add_run("Single execution blueprint for engineering interns building TMT's internal operations platform from lead acquisition and corporate ATS recruitment through training, company internship, work execution, asset management, leave workflows, communications, finance, document generation, exit offboarding, certificates, reporting, and audit.")
    r4_body.font.size = Pt(10)

    p5 = doc.add_paragraph(style='Normal')
    p5.paragraph_format.space_after = Pt(12)
    r5 = p5.add_run("INTERNAL • IMPLEMENTATION USE")
    r5.bold = True
    r5.font.size = Pt(9)
    r5.font.color.rgb = RGBColor(0x82, 0x82, 0x82)

def add_heading_1(text):
    p = doc.add_paragraph(style='Heading 1')
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)

def add_heading_2(text):
    p = doc.add_paragraph(style='Heading 2')
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(11.5)
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)

def add_heading_3(text):
    p = doc.add_paragraph(style='Heading 3')
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)

def add_body_paragraph(text):
    if not text.strip():
        return
    p = doc.add_paragraph(style='Normal')
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    # Parse bold formatting
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2])
            r.bold = True
        else:
            r = p.add_run(part)
        r.font.size = Pt(9.5)

def add_bullet_item(text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2])
            r.bold = True
        else:
            r = p.add_run(part)
        r.font.size = Pt(9.5)

def add_numbered_item(text):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2])
            r.bold = True
        else:
            r = p.add_run(part)
        r.font.size = Pt(9.5)

def add_callout_block(text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    set_cell_background(cell, "EAF2F8")
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    cell.width = Inches(7.1)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    clean_text = text.replace("> ", "").strip()
    parts = re.split(r'(\*\*.*?\*\*)', clean_text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2])
            r.bold = True
        else:
            r = p.add_run(part)
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0x17, 0x36, 0x5D)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)

def add_code_block(text):
    p = doc.add_paragraph(style='CodeBlock')
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

def add_table_from_md(table_lines):
    rows = []
    for line in table_lines:
        line = line.strip()
        if not line.startswith('|') or not line.endswith('|'):
            continue
        # Split cells
        cells = [c.strip() for c in line[1:-1].split('|')]
        # Skip separator row e.g. |---|---|
        if all(re.match(r'^:?-+:?$', c) for c in cells):
            continue
        rows.append(cells)
    
    if len(rows) < 2:
        return
    
    headers = rows[0]
    data = rows[1:]
    col_count = len(headers)
    
    tbl = doc.add_table(rows=len(data) + 1, cols=col_count)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Format headers
    h_row = tbl.rows[0]
    for idx in range(col_count):
        cell = h_row.cells[idx]
        cell.text = headers[idx].replace('**', '')
        set_cell_background(cell, "1F4E78")
        set_cell_margins(cell, top=80, bottom=80, start=100, end=100)
        set_cell_valign(cell, "center")
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(9)
                
    # Format data rows
    for r_idx, r_data in enumerate(data):
        row = tbl.rows[r_idx + 1]
        bg = "F4F7FA" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx in range(col_count):
            cell = row.cells[c_idx]
            val = r_data[c_idx] if c_idx < len(r_data) else ""
            cell.text = ""
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, start=90, end=90)
            set_cell_valign(cell, "top")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            parts = re.split(r'(\*\*.*?\*\*)', val)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    r = p.add_run(part)
                r.font.size = Pt(8.5)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

# Parse Markdown into blocks
add_title_block()

lines = md_content.split('\n')
i = 0
in_toc = False

while i < len(lines):
    line = lines[i].rstrip()
    
    # Skip preamble lines that are already in title block
    if line.startswith('# TECHMECHA') or line.startswith('# TMT OPERATIONS') or line.startswith('## Intern Implementation') or line.startswith('### Version 1.1') or line.startswith('**Purpose:**') or line.startswith('**INTERNAL'):
        i += 1
        continue
    if line.startswith('---'):
        i += 1
        continue
        
    # Check Table of Contents section
    if line.startswith('### Contents'):
        add_heading_1("Contents")
        in_toc = True
        i += 1
        continue
        
    if in_toc:
        if line.startswith('## 1. Product Goal'):
            in_toc = False
            # proceed to parse section 1
        elif line.strip() and line[0].isdigit():
            # e.g. "1. [Product Goal and Scope](#1-product-goal-and-scope)"
            m = re.match(r'^(\d+)\.\s+\[(.*?)\]', line)
            if m:
                add_numbered_item(f"{m.group(1)}. {m.group(2)}")
            else:
                add_numbered_item(line)
            i += 1
            continue
        else:
            i += 1
            continue
            
    # Markdown Table
    if line.startswith('|') and '|' in line[1:]:
        table_lines = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            table_lines.append(lines[i].strip())
            i += 1
        add_table_from_md(table_lines)
        continue
        
    # Codeblock
    if line.startswith('```'):
        code_lines = []
        i += 1
        while i < len(lines) and not lines[i].startswith('```'):
            code_lines.append(lines[i])
            i += 1
        i += 1 # skip closing ```
        add_code_block('\n'.join(code_lines))
        continue
        
    # Callout / Blockquote
    if line.startswith('>'):
        callout_lines = []
        while i < len(lines) and lines[i].startswith('>'):
            callout_lines.append(lines[i][1:].strip())
            i += 1
        add_callout_block(' '.join(callout_lines))
        continue
        
    # Headings
    if line.startswith('## '):
        add_heading_1(line[3:].strip())
        i += 1
        continue
    elif line.startswith('### '):
        add_heading_2(line[4:].strip())
        i += 1
        continue
    elif line.startswith('#### '):
        add_heading_3(line[5:].strip())
        i += 1
        continue
        
    # Lists
    if line.startswith('- ') or line.startswith('* '):
        add_bullet_item(line[2:].strip())
        i += 1
        continue
    elif re.match(r'^\d+\.\s+', line):
        clean_num = re.sub(r'^\d+\.\s+', '', line)
        add_numbered_item(clean_num)
        i += 1
        continue
        
    # Regular paragraph
    if line.strip():
        add_body_paragraph(line.strip())
        
    i += 1

doc.save(OUTPUT_PATH)
print(f"Successfully generated enhanced DOCX: {OUTPUT_PATH}")
