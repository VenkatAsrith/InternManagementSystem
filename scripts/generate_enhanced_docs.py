# -*- coding: utf-8 -*-
"""
TMT Operations Platform - Master Documentation & DOCX Generator
Version 1.1 Enterprise Enhanced | 17 September 2026
Integrates corporate ATS recruitment intake, asset management, leave management,
10-stage exit & offboarding lifecycle, TMT Letter Studio vector engine, and 14 enterprise roles
while strictly preserving the canonical 28-section flow and domain principles.
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOCX_TEMPLATE = "TMT_Operations_Platform_Intern_Implementation_Ready_Documentation_v1.0_ORIGINAL.docx"
DOCX_OUTPUT = "TMT_Operations_Platform_Intern_Implementation_Ready_Documentation_v1.0 (1).docx"
MD_OUTPUT = os.path.join("docs", "TMT_OPERATIONS_PLATFORM_V1.1_MASTER.md")

# XML helper functions for Word styling
def set_cell_background(cell, fill_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=80, bottom=80, start=100, end=100):
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{start}" w:type="dxa"/><w:right w:w="{end}" w:type="dxa"/></w:tcMar>')
    cell._tc.get_or_add_tcPr().append(tcMar)

def set_cell_valign(cell, valign="center"):
    valign_el = parse_xml(f'<w:vAlign {nsdecls("w")} w:val="{valign}"/>')
    cell._tc.get_or_add_tcPr().append(valign_el)

def style_table_header(row, col_widths=None):
    for idx, cell in enumerate(row.cells):
        set_cell_background(cell, "1F4E78")
        set_cell_margins(cell, top=100, bottom=100, start=120, end=120)
        set_cell_valign(cell, "center")
        if col_widths and idx < len(col_widths):
            cell.width = col_widths[idx]
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(9.5)

def style_table_body(table, col_widths=None):
    for r_idx, row in enumerate(table.rows[1:]):
        bg = "F4F7FA" if r_idx % 2 == 1 else "FFFFFF"
        for idx, cell in enumerate(row.cells):
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, start=100, end=100)
            set_cell_valign(cell, "top")
            if col_widths and idx < len(col_widths):
                cell.width = col_widths[idx]
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)

def add_styled_table(doc, headers, data, col_widths=None):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    header_row = tbl.rows[0]
    for idx, h in enumerate(headers):
        header_row.cells[idx].text = str(h)
    style_table_header(header_row, col_widths)
    
    for r_idx, row_data in enumerate(data):
        row = tbl.rows[r_idx + 1]
        for c_idx, val in enumerate(row_data):
            row.cells[c_idx].text = str(val)
    style_table_body(tbl, col_widths)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    return tbl

def add_callout(doc, text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    set_cell_background(cell, "EAF2F8")
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    cell.width = Inches(7.1)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0x17, 0x36, 0x5D)
    run.bold = False
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_codeblock(doc, code_text):
    p = doc.add_paragraph(style='CodeBlock')
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.font.size = Pt(9.5)

def add_numbered(doc, text):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.font.size = Pt(9.5)

def add_heading_1(doc, text):
    p = doc.add_paragraph(style='Heading 1')
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
    r.bold = True

def add_heading_2(doc, text):
    p = doc.add_paragraph(style='Heading 2')
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
    r.bold = True

def add_body(doc, text):
    p = doc.add_paragraph(style='Normal')
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.font.size = Pt(9.5)

print("Helper functions initialized.")
